"""Build the implementation-based project explanation as a Gill Sans MT DOCX."""
from pathlib import Path
import ast
import gzip
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Project_Complete_Explanation.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = 'Gill Sans MT'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Gill Sans MT')
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:cs'), 'Gill Sans MT')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
for name, size in [('Title', 27), ('Heading 1', 19), ('Heading 2', 14), ('Heading 3', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor.from_string('17365D')
header = sec.header.paragraphs[0]
header.text = 'VEHICULAR EMERGENCY ALERTS | PROJECT EXPLANATION'
header.style = doc.styles['Caption']
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Page ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text):
    doc.add_paragraph(text)

def h(text, level=1):
    doc.add_heading(text, level)

def table(headings, rows):
    t = doc.add_table(rows=1, cols=len(headings)); t.style = 'Light Shading Accent 1'
    for c, text in zip(t.rows[0].cells, headings): c.text = str(text)
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    for row in t.rows:
        for c in row.cells:
            for para in c.paragraphs:
                for r in para.runs: r.font.size = Pt(9)
    trpr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph()

def formula(expr, explanation):
    para = doc.add_paragraph()
    para.add_run(expr).bold = True
    p(explanation)

def figure(filename, caption):
    path = ROOT / 'results/urban/figures' / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.35))
        doc.add_paragraph(caption, 'Caption')

doc.add_heading('Complete Project Explanation', 0)
p('Trust-Aware Cluster-Based Broadcast Suppression for Real-Time Vehicular Emergency Alerts')
p('Algorithms • formulas • purpose of each component • complete SUMO workflow • security • evaluation • operating instructions')
p('Prepared from the local implementation on 16 October 2026. Written in English for project study, demonstration and viva preparation. All editable document text uses Gill Sans MT; existing screenshots retain their original image lettering.')
p('This document describes the code as implemented. Numerical examples are teaching examples unless explicitly labelled as recorded simulation evidence. Literature comments in source files are not independently verified citations.')
doc.add_page_break()
h('Contents')
sections = ['Project purpose and complete flow', 'Technology stack and file responsibilities', 'SUMO road network and vehicle generation', 'Dataset and incident generation', 'Radio geometry and channel simulation', 'Watchdog observations and Beta trust', 'DBSCAN clustering', 'Cluster representatives and bridge roles', 'Voting and event corroboration', 'BLS signatures and certificate validation', 'Forwarding protocols and suppression', 'Cluster-head shortest-path routing', 'RSU validation and duplicate handling', 'Storage, trust updates, cloud and dispatch', 'SUMO demonstration and visual meanings', 'Metrics and statistical formulas', 'Experiments, ablations and forged alerts', 'Web application and AI explanations', 'Commands and practical usage', 'Recorded example, limitations and viva answers', 'Configuration reference and source index']
for i, title in enumerate(sections, 1): p(f'{i}. {title}')
doc.add_page_break()

h('1. Project purpose and complete flow')
p('The project investigates how vehicles can deliver an emergency alert to nearby vehicles and roadside infrastructure while reducing unnecessary repeated broadcasts. Blind flooding makes almost every honest receiver transmit the same message. In dense traffic this creates contention, duplicate receptions and collisions: the broadcast-storm problem.')
p('The proposed method combines local density clustering, reputation-based relay priority, distance-based waiting, random delay, and a neighbourhood coverage test. A separate security stage requires distinct endorsements and a real aggregate signature before an alert is accepted. The project also compares a route restricted to trusted cluster heads for delivery to an RSU.')
p('End-to-end flow: SUMO traffic → position/speed snapshot → radio neighbour graph → trust and DBSCAN roles → staged accident → signed proposal → delivered votes → aggregate certificate → selected forwarding rule → RSU validation → recorded results and local infrastructure storage.')
p('SUMO supplies mobility and the graphical road scene. Python supplies the communication model, protocol decisions, trust calculations, cryptography, measurements and storage. This is a coupled research simulation, rather than a native SUMO wireless-network feature or a deployed emergency service.')

h('2. Technology stack and file responsibilities')
table(['Tool/library', 'Purpose and how it is used'], [
('Python', 'Coordinates scenario generation, algorithms, experiments, reports, web server and infrastructure adapters.'),
('SUMO / eclipse-sumo', 'Microscopic vehicle mobility, passenger-car behaviour, roads, junctions and GUI playback.'),
('TraCI', 'Controls SUMO-GUI vehicles, speeds, colours, polygons and presentation state.'),
('sumolib / netconvert', 'SUMO support utilities and conversion of node/edge XML into the drivable network.'),
('NumPy', 'Coordinates, adjacency matrices, masks, random streams and vectorised calculations.'),
('scikit-learn', 'DBSCAN implementation for density-based clustering of vehicle coordinates.'),
('SciPy', 'Connected components for reachable vehicles and Student-t confidence intervals.'),
('py_ecc', 'Actual BLS12-381 signing, public-key derivation, aggregation and signature verification.'),
('Matplotlib', 'Static figures, comparisons, maps, timelines and animation assets.'),
('SQLite / JSON / gzip / CSV', 'Durable local backhaul, detailed compressed traces, configuration and summary datasets.'),
('urllib / HTTP server', 'Optional HTTPS APIs and local browser interface without an additional web framework.'),
('Groq API', 'Optional language-model explanations of recorded evidence; does not decide protocol outcomes.'),
('pytest', 'Validation suite covering channel, security, protocols, RSU, pipeline and user-facing explanations.')])
table(['Source file', 'Responsibility'], [
('config.py', 'Central parameters, protocol names, experimental variants and selected network.'),
('scenario.py', 'Routes, deterministic seeds, SUMO execution, FCD parsing, watchdog observations and accident packet.'),
('clustering.py', 'Beta trust, DBSCAN, greedy head election, bridge/member labels and priority tiers.'),
('corroboration.py', 'Keys, signed event bytes, committee planning, aggregate certificates and acceptance.'),
('channel.py', 'Geometry, frame airtime, CSMA/CA, interference and event-driven dissemination.'),
('protocols.py', 'Proposal/vote handling, flooding, slotted persistence, proposed suppression and CH routing.'),
('rsu.py', 'Packet plausibility checks and once-per-event registration.'),
('infrastructure.py', 'Dijkstra routing, certificate-list processing, trust updates, SQLite, export and dispatch.'),
('experiment.py', 'Shared scenario comparisons, forged-alert cases, raw traces and aggregated experiments.'),
('metrics.py', 'Definitions, measurements, consistency checks and confidence intervals.'),
('demo.py / presentation.py', 'TraCI replay, cluster display, overlays and infrastructure status.'),
('figures.py / report.py', 'Plots, explanations, reports and static dashboard.'),
('narrator.py / web_results.py', 'Optional AI narration and multilingual saved-result application.'),
('run.py', 'Command-line entry point for all supported operations.'),
('network/build_urban.py', 'Builds the urban roads and decorative scenery.'),
('validation/', 'Executable validation tests; read these alongside the relevant implementation.')])

h('3. SUMO road network and vehicle generation')
p('The selected network is network/urban.net.xml. Its builder defines a 5 × 5 layout with junction IDs A0 through E4. The nominal x coordinates are 0, 160, 390, 610 and 820 metres; y coordinates are 0, 170, 370, 590 and 800 metres. Interior junctions use traffic lights and border junctions use priority control. Adjacent junctions are connected in both directions.')
p('Most generated edges have two lanes; the central horizontal Garden Boulevard uses three lanes in each generated direction. The configured edge speed is 13.89 m/s, approximately 50 km/h. Buildings, trees, ponds and landscaped blocks are display objects. They do not block radio signals in this unit-disk model.')
p('Two RSUs are placed at actual junction coordinates read from the selected network: RSU-1 at B3 and RSU-2 at D1. This runtime lookup supersedes the preliminary numeric coordinate dictionary near the beginning of config.py. Areas outside their direct range make multi-hop vehicle relaying relevant.')
p('Each vehicle receives a random-walk route of up to 30 edges. The generator avoids immediate U-turns when other choices exist. Departure times are uniform over the first 30 seconds and sorted before writing routes. SUMO uses its default passenger-car type, including the default Krauss car-following behaviour. Movement is sampled every one second through floating-car-data XML.')
formula('seed(d, r) = 10000 + 100d + r', 'd is the requested vehicle population and r is the run index. For d = 200 and r = 2, the seed is 30002. It fixes scenario randomness, enabling repeatable mobility and incident selection. Protocol-specific random streams still differ between schemes.')

h('4. Dataset and incident generation')
p('The dataset is generated by simulation, not downloaded from a real accident dataset. Each FCD timestep yields vehicle ID, x/y position in metres, speed in m/s, angle and lane. The saved incident dataset adds trust, malicious status, neighbour count, cluster ID and role. The requested population may differ from vehicles active at the snapshot if SUMO insertion or route completion changes availability; metrics use the active set.')
p('An incident time is sampled as an integer between 60 and 120 seconds, inclusive. At that time a randomly selected active non-malicious vehicle is chosen as the genuine incident source. Its position and lane define the staged accident. This is controlled event injection: the program does not infer a crash from impact forces or camera/sensor data.')
table(['Alert field', 'Meaning'], [('event_id', 'EVT-{seed}-{vehicle}-{sequence}: simulation event identity.'), ('message_type', 'ACCIDENT, currently the recognised event category.'), ('source_vehicle / sequence_number', 'Original sender identity and event sequence.'), ('incident_position', 'Snapshot x/y coordinates, rounded to two decimal places.'), ('incident_time_s / created_s', 'Detection and creation times; equal at event generation.'), ('hop_limit', 'Maximum permitted forwarding depth, normally 10.')])
p('One genuine incident is reused across forwarding schemes. Forgery tests create a separate false event on the same traffic snapshot using a malicious source. Each individual dissemination simulation contains one alert. Event IDs support deduplication; IDs alone do not authenticate a message.')

h('5. Radio geometry and channel simulation')
formula('d(i,j) = sqrt((x_i − x_j)^2 + (y_i − y_j)^2)', 'Euclidean distance between two radio nodes. Nodes include active vehicles followed by RSUs.')
formula('A(i,j) = 1 when i ≠ j and d(i,j) ≤ R; otherwise 0', 'A is the symmetric adjacency matrix; R = 300 m by default. The same range is used for decoding, carrier sensing and interference. Self-links are removed. The matrix also acts as an idealised beacon neighbour table; beacon traffic itself is not transmitted on the simulated channel.')
p('The channel approximates 802.11p broadcast CSMA/CA using 13-microsecond slots. A sender waits for an idle arbitration interval, chooses random backoff from 0 to 3 slots, freezes backoff when the medium becomes busy, and transmits when the countdown ends. Broadcast has no ACK or retry and no increasing contention window. Simultaneous countdown completion and hidden terminals can cause collisions.')
formula('AIFS = SIFS + AIFSN × slot = 32 + 2 × 13 = 58 microseconds', 'The implementation rounds AIFS upward to 5 radio slots, effectively 65 microseconds. Discrete-time rounding matters when interpreting short delays.')
formula('T_air = 40 + 8 × ceil((16 + 6 + 8(L + 38))/48) microseconds', 'L is the supplied payload size in bytes, including any certificate material passed by the protocol. The expression models a 40-microsecond preamble and 48 data bits per 8-microsecond OFDM symbol at the fixed 6-Mbit/s rate. MAC/LLC/FCS overhead totals 38 bytes.')
formula('frame_slots = ceil(T_air / 13); time_ms = slots × 13/1000', 'For a bare 300-byte alert, T_air = 496 microseconds and the channel occupies 39 slots, or 507 microseconds. Signed proposals, votes and certified alerts have different supplied payload sizes and therefore different airtimes.')
p('A candidate reception is lost if another transmitter overlaps the frame and is in range of that receiver, or the receiver is itself transmitting. This captures interference and half-duplex operation. No capture effect, fading, buildings, propagation model or background load is simulated. The dissemination loop normally allows up to 2,000 ms and jumps between events when no sender is contending.')
p('The airtime equation currently hard-codes the 6-Mbit/s symbol structure. Changing DATA_RATE_MBPS alone does not implement a new physical rate. This distinction matters when designing sensitivity experiments.')

h('6. Watchdog observations and Beta trust')
p('Trust estimates past forwarding reliability. At each one-second mobility sample, a vehicle with at least one neighbour receives a simulated watchdog observation. Malicious vehicles count as failures. Honest forwarding can also be missed with probability 0.10. These are synthetic observations tied to scenario ground truth, not packet-sniffing evidence from a deployed watchdog.')
formula('T_i = (s_i + 1)/(s_i + f_i + 2)', 's_i and f_i are observed successes and failures. The formula is the posterior mean of a Beta distribution with uniform Beta(1,1) prior. An unobserved vehicle starts at 0.5. The trust gate is T_i ≥ 0.7.')
formula('Example: s = 18, f = 2 → T = 19/22 ≈ 0.864', 'This vehicle passes the 0.7 threshold. A vehicle with s = 0 and f = 20 has T = 1/22 ≈ 0.045 and fails it. Missing an observation is different from observing a failure: no-neighbour timesteps do not add evidence.')
p('Trust affects representative election, early forwarding tiers, committee eligibility and certificate endorsement counting. A low-trust receiver can still receive a valid alert and may use the later fallback tier in the proposed dissemination scheme. Reputation is shared globally in this simulation; distributed trust exchange and adversarial trust manipulation are outside the model.')

h('7. DBSCAN clustering')
p('DBSCAN means Density-Based Spatial Clustering of Applications with Noise. It groups spatially dense vehicles without specifying a fixed cluster count. The implementation fits DBSCAN to x/y positions only; speed is used later to break head-election ties. The neighbourhood radius is 40 m and min_samples is 3, including the point itself under scikit-learn semantics.')
formula('N_eps(i) = {j : d(i,j) ≤ 40 m}; core point when |N_eps(i)| ≥ 3', 'A core point lies in a dense neighbourhood. Density-connected core points form a cluster; nearby border points can join. Vehicles not assigned to a cluster receive label −1, called noise.')
p('The 40-m clustering radius and 300-m communication radius serve different purposes. A vehicle can be a DBSCAN noise point yet still communicate with several radio neighbours. Noise therefore does not mean disconnected. During the mobility demonstration clusters are recomputed each second; the subsequent dissemination uses the incident snapshot.')

h('8. Cluster representatives and bridge roles')
p('Within each DBSCAN group, trusted candidates are selected by greedy coverage. Start with every group member uncovered. Choose the trusted vehicle covering the most uncovered group members within radio range, including itself. Mark those members covered and repeat until all are covered or no useful trusted candidate remains. A group may need multiple representatives, or may have uncovered members when no trusted candidate can help.')
formula('coverage_gain(i) = |(N_radio(i) ∪ {i}) ∩ uncovered_members|', 'Maximise this value to reduce the number of representative relays needed for the group. This is a greedy set-cover heuristic and does not guarantee the mathematical optimum.')
formula('relative_speed(i) = mean over group neighbours j of |speed_j − speed_i|', 'When coverage gains tie, choose the lower mean relative speed. Similar speeds suggest more stable neighbourhood relationships. With no group radio neighbours the implementation uses zero for this tie-break value.')
table(['Role/tier', 'Meaning'], [('CH', 'Trusted vehicle elected as a cluster representative.'), ('MEMBER', 'Clustered vehicle that was not elected.'), ('BRIDGE', 'DBSCAN noise vehicle; a potential relay across sparse areas.'), ('Tier 0', 'CH or BRIDGE whose trust passes the threshold; decides earlier.'), ('Tier 1', 'All other vehicles; later forwarding opportunity.')])
p('Bridge is a forwarding role, not a graph-theoretic articulation-point test. A CH is not automatically the only vehicle allowed to forward in the proposed broadcast scheme: later-tier vehicles remain available if useful coverage is missing.')

h('9. Voting and event corroboration')
p('A genuine source first broadcasts a signed PROPOSAL. Up to five nearest in-range reputable neighbours form the intended committee. Honest committee members observe the staged event as true and send VOTE frames. Malicious members stay silent for genuine incidents. Votes must actually survive radio delivery to the source; planned voters and delivered votes are different quantities.')
p('After at least three delivered endorsements, the source schedules certificate release with the configured verification cost. The certificate contains the source signature plus received voter signatures. The origin source does not count toward the three required endorsers. A committee of fewer than three eligible delivered voters cannot release the protected alert.')
formula('Origin acceptance: distinct valid endorsers with trust ≥ 0.7 must number at least 3', 'This is the rule implemented by corroboration.accept. “3 of 5” is a small corroboration quorum; it should not be described as a complete Byzantine consensus protocol with formal network-wide fault guarantees.')
p('Voters incur nominal verification plus signing delays, with rank-based 0.5-ms offsets. Their frames still contend on the shared channel; offsets do not remove possible interference or failed deliveries. Nominal signing and verification charges are 1.4 ms and 4.3 ms. These are model parameters, separate from Python wall-clock execution time.')
p('In the forged-alert scenario honest vehicles observe no real incident and do not endorse. Nearby malicious vehicles can collude and produce genuine signatures over false content. The reputation gate then determines which endorsements count. Cryptographic validity alone cannot establish event truth.')

h('10. BLS signatures and certificate validation')
p('BLS means Boneh–Lynn–Shacham. The project uses py_ecc G2ProofOfPossession over BLS12-381. Each simulation vehicle gets a deterministic secret key derived from SHA-256 of its ID with a simulation prefix, reduced into the curve scalar range. This is repeatable test identity generation, not secure deployment key management.')
p('The signed byte string concatenates event ID, message type, source vehicle, sequence, incident x/y rounded to two decimals and incident time to three decimals. The function is called digest, but returns canonical event bytes rather than a separately computed SHA-256 event digest. BLS internally hashes the message to its signing group. Hop count is excluded so relaying does not require changing the event signature.')
formula('Conceptually: PK_i = sk_i × G; sigma_i = sk_i × H(message); Sigma = sum_i sigma_i', 'This additive notation explains the relationship. H maps the event bytes to the appropriate curve group and the real operations are performed by py_ecc. A pairing-based aggregate check binds the common message to all listed public keys.')
p('A certificate stores signer IDs and the aggregate signature encoded as hexadecimal. The aggregate is 96 bytes before hex encoding, independent of signer count; signer IDs still add size. A receiver rejects duplicate signer identities, an incorrect source identity, invalid aggregate mathematics, or fewer than three qualifying endorsers. Verification uses FastAggregateVerify for one common message.')
p('The acceptance routine does not independently enforce that all signers belonged to the originally planned nearest-five committee; normal protocol construction supplies that selection. The implementation also does not provision a deployed PKI or explicit per-vehicle proof-of-possession registration. These boundaries belong in any security claim.')
p('The simulator shares a computed certificate verdict and adds nominal receiver verification latency. It does not benchmark fresh independent pairing operations on every receiver CPU. Caching further separates runtime performance from the charged cryptographic model.')

h('11. Forwarding protocols and suppression')
p('All protected baselines share the same origin voting and BLS mechanism. Forwarding begins only after certificate release. Each vehicle keeps per-event accepted/seen flags, reception time, previous sender, hop count, heard senders and a forwarding decision. Malicious vehicles receive but never relay. RSUs are receive-only endpoints. A receiver at the hop limit does not relay further.')
h('11.1 Blind flooding', 2)
p('Every honest receiver of its first accepted certified copy queues a rebroadcast immediately. The MAC can still delay or lose it. Later copies do not cause another rebroadcast. Flooding is a useful delivery baseline but can produce many redundant frames and collisions.')
h('11.2 Slotted 1-persistence', 2)
formula('S(D) = clamp(Ns − ceil(Ns × min(D,R)/R), 0, Ns−1)', 'D is distance from the first sender, R = 300 m and Ns = 5. The receiver waits S(D) × 5 ms. Greater distance yields an earlier slot, aiming to make more progress per transmission.')
p('For D = 270 m the slot is 0; for 150 m it is 2; for 30 m it is 4. If a duplicate arrives before its timer expires, the vehicle suppresses its pending protocol transmission. Otherwise it forwards. The use of duplicate hearing alone can leave coverage gaps.')
h('11.3 Proposed trust-aware suppression', 2)
formula('wait_i = (tier_i × Ns + S(D_i)) × tau + RAD_i', 'Ns = 5, tau = 5 ms and RAD is sampled on the radio-slot grid over approximately [0,5 ms). Tier 0 receives the first block of distance slots; tier 1 waits an additional 25 ms. Exact timings are quantised using ms_to_slots.')
p('For tier 0 and D = 150 m, the nominal wait is 10 ms plus random delay. For tier 1 at the same distance it is 35 ms plus random delay. The random component reduces simultaneous decisions among vehicles in the same priority/distance group.')
formula('covered_i = union over heard senders h of (N_radio(h) ∪ {h})', 'The receiver estimates which nodes are already geometrically covered by senders whose certified copies it decoded.')
formula('uncovered_i = N_radio(i) minus covered_i minus {i}', 'When the timer expires, forward if this set is nonempty; otherwise suppress. RSUs count as neighbours, so an uncovered RSU can motivate a relay. Geometric coverage estimates do not prove every neighbour actually decoded a packet, particularly when collisions occur.')
p('Example: a candidate has neighbours A, B and an RSU. If previous heard senders cover A and B but not the RSU, it forwards. If every neighbour is already covered, another broadcast is unnecessary under this estimate. Savings must be assessed alongside delivery and delay, since longer waits and imperfect coverage estimates create trade-offs.')

h('12. Cluster-head shortest-path routing')
p('The ch_route alternative builds a weighted graph containing the source, trusted non-malicious elected CHs and the RSUs. Only in-range edges exist. Bridge vehicles are not eligible intermediates in this route. Dijkstra’s algorithm finds the minimum sum of Euclidean link distances to an available RSU, subject to the hop limit.')
formula('cost(path) = sum over links (u,v) in path of d(u,v)', 'The route cost is metres. It does not optimise measured delay, channel congestion or probability of successful delivery. The implementation tracks best costs by node and hop count to enforce the hop budget.')
p('At each intermediate head, a local signed proposal triggers radio-delivered votes from trusted honest same-cluster neighbours. The head counts its own distinct endorsement and requires three distinct endorsements in total before forwarding. This local rule differs from origin corroboration, where the source signature is separate from three neighbour endorsements.')
p('A geometric route can exist but fail because local quorum is insufficient or vote frames are lost. If no eligible CH route exists, there is no hidden bridge/flooding fallback. Non-route vehicles can receive the alert but remain silent. Consequently CH-route delivery should be judged mainly by RSU success, hops, transmissions and delay rather than by vehicle-wide reachability.')
p('The original event certificate stays associated with dissemination, while local round certificates are also recorded as consensus evidence. Batch processing later verifies certificate entries separately and deduplicates accepted event IDs; it is not a multi-message cryptographic pairing speedup. infrastructure.cluster_votes is a separate helper; actual CH-route local rounds are radio-modelled in protocols.CHRoute.')

h('13. RSU validation and duplicate handling')
p('RSUs receive but never relay, so they do not create artificial paths between vehicle components. rsu.py validates required packet fields, recognised ACCIDENT type, position within the network bounding box, creation/incident time order, freshness within 600 seconds, hop limit and the supplied certificate verdict. The position check is a bounding-box plausibility check, not map matching to the exact road surface.')
p('All RSUs contribute to a common per-event registry. The earliest valid received copy creates the event record; later copies, including those at another RSU, count as duplicates. The record includes first RSU, sender, delay, hop count and transmission count before receipt. This prevents repeated registration of the same event within the processed run.')
p('No-security ablations skip the certificate requirement intentionally. Plausibility checks cannot prove a crash happened. Real sensor validation, deployed identity management and actual traffic-management backhaul remain outside this radio model.')

h('14. Storage, trust updates, cloud and dispatch')
p('The experiment stores incident datasets, CH-route accepted events and trust observations in results/urban/infrastructure/backhaul.sqlite3. Events and dispatch use a composite run_id/event_id key; trust uses run_id/vehicle_id. The database also stores one current dataset row per run ID. Re-running a seed can replace those database rows, while timestamped raw result files preserve separate invocation history.')
p('An accepted CH-route event adds a QUEUED_LOCAL outbox entry. This is a durable local record. It does not contact an ambulance or emergency operator during simulation. Infrastructure persistence follows the CH-route outcome in run_one, which can differ from the proposed broadcast scheme’s RSU outcome.')
formula('T_after = (20 × T_before + 1)/21 for an observed successful forward', 'This is an explicit effective prior of weight 20, rather than a reconstruction of original Beta success/failure counts. For T_before = 0.8, the updated value is about 0.809524. If there is no accepted event or no observed transmission for the vehicle, trust stays unchanged. Silence is not penalised by this update.')
p('python run.py sync always creates a local JSON export. It uploads only when IOV_CLOUD_URL provides an HTTPS ingestion endpoint; IOV_CLOUD_TOKEN optionally supplies bearer authentication. The payload contains events, trust updates, dispatch state and vehicle datasets.')
p('python run.py dispatch sends pending accepted records only when IOV_EMERGENCY_URL is configured, optionally with IOV_EMERGENCY_TOKEN. Successful HTTP 2xx delivery marks the row DELIVERED_REMOTE. Requests include Idempotency-Key headers, which the receiving API must honour to prevent duplicate processing across retries. No cloud account, physical LAN model or real dispatch endpoint is bundled.')

h('15. SUMO demonstration and visual meanings')
p('The GUI is a presentation of a previously computed radio trace. It does not decide relay outcomes as the animation runs. The code recomputes live DBSCAN roles during the pre-incident mobility view and checks that its incident snapshot matches the experiment. Replay then slows millisecond events so a person can see them.')
table(['Stage', 'What happens and what it demonstrates'], [('Normal traffic', 'Vehicles move under SUMO rules; groups and representatives change with position.'), ('Incident', 'The selected vehicle stops and turns red; its staged event becomes the alert source.'), ('Voting', 'Proposal/vote activity is displayed before certified alert release.'), ('Radio replay', 'Vehicles freeze at the incident snapshot; rings and colours replay logged forwarding decisions.'), ('Comparison', 'Flooding, proposed and CH-route views reuse the same incident for interpretation.'), ('RSU', 'The display reports actual recorded acceptance, rejection or non-delivery.'), ('Forgery', 'A separate false alert illustrates the selected protection outcome; success/rejection depends on recorded evidence.'), ('Finish', 'Other traffic resumes around the stopped incident vehicle.')])
table(['Display', 'Meaning'], [('Cluster colours', 'One palette colour per DBSCAN group; white for isolated points in the clustering view.'), ('Black ring', 'Elected cluster representative.'), ('Yellow', 'Vehicle has not yet accepted the alert.'), ('Violet', 'Committee voting activity.'), ('Blue', 'Vehicle accepted the alert.'), ('Orange / transmission ring', 'Rebroadcast and sender communication range.'), ('Grey', 'Forwarding suppressed.'), ('Purple drop', 'Misbehaving receiver that does not relay.'), ('White in alert view', 'Rejected alert; context differs from white isolated points in the clustering view.'), ('Red source', 'Incident or selected alert-origin vehicle.'), ('Cyan / green / red RSU', 'Listening / validated / rejected infrastructure state.')])
p('Traffic freezes because mobility is sampled in one-second steps while communication takes milliseconds; the model intentionally uses a fixed incident topology. Animation pace is a wall-clock presentation delay, not measured emergency latency. Map scenery and ring overlays help explanation but do not add physical radio obstructions.')
figure('sumo_1_clusters.png', 'Existing project screenshot: SUMO clustering view. Screenshot labels retain their original rendering.')
figure('sumo_2_accident.png', 'Existing project screenshot: staged incident view; this is illustrative, not a new experiment run.')

h('16. Metrics and statistical formulas')
p('Let e be vehicles in the source’s connected component of the vehicle-only radio graph, r the vehicles holding the accepted alert including the source, t the number of certified ALERT frames including the source, and H the honest non-source accepted receivers. The component denominator is shared across schemes; RSUs are excluded as connectivity relays.')
table(['Metric/formula', 'Interpretation'], [('Reachability = r/e', 'Fraction of theoretically connected vehicles reached. Undefined for the implementation’s isolated-source case e ≤ 1.'), ('Saved rebroadcast = (r−t)/r', 'Fraction of accepted receivers saved from an alert transmission; undefined when no alert is released.'), ('Suppression = 1 − (t−1)/H', 'Share inferred silent among honest non-source receivers; undefined when H = 0.'), ('Total transmissions = len(channel log)', 'All on-air frames: proposals, votes, certified alerts and CH-local exchanges.'), ('Redundant alert transmissions', 'Certified-alert frames delivering to no new vehicle or RSU.'), ('Duplicate receptions', 'Extra accepted copies at vehicles; RSU duplicate accounting is separated.'), ('Collision ratio = lost/(successful + lost)', 'Losses count receiver opportunities, not just collided transmitted frames.'), ('t90', 'Acceptance time of the ceil(0.9e)-th reached vehicle; undefined if the target is not reached.'), ('Mean delay', 'Mean accepted-reception time excluding the source.'), ('Last delay', 'Latest accepted-reception time, including source as time zero.'), ('CBR over 100 ms', 'Per-vehicle busy-time fraction in the first 100 ms, averaged over the reachable component.'), ('Airtime', 'Mean cumulative sensed busy time across reachable vehicles; overlapping time is counted once per vehicle.'), ('RSU received', '1 if registered, 0 if reachable but not registered, undefined if no RSU was reachable.'), ('RSU delay / hops', 'Incident-to-first-registration time and hop count of the registered copy.'), ('RSU relay-only metrics', 'Computed only for source incidents outside direct RSU range that have an RSU-reachable vehicle component.'), ('Quorum delay', 'Incident-to-certified-alert release time, before later dissemination.'), ('Forgery accepted share', 'Honest non-source vehicles accepting a false alert divided by honest non-source vehicles reachable in principle.')])
formula('CBR_i = busy_slots_i(first 100 ms)/slots_in_100ms', 'The reported cbr_100ms averages this over the source-connected vehicles. cbr_peak is the largest such per-vehicle fraction, rather than a sliding-window temporal maximum. Unused remaining time in the fixed 100-ms window contributes idle time.')
formula('mean = sum(x_i)/n; s = sqrt(sum((x_i−mean)^2)/(n−1))', 'Summaries remove undefined NaN values separately for each metric, so contributing sample counts can differ between columns.')
formula('95% confidence interval = mean ± t(0.975, n−1) × s/sqrt(n)', 'Student-t intervals express uncertainty of the estimated mean under the sampling assumptions. The implementation emits zero half-width for n = 1 as a formatting convention; one run does not establish precision. At small n, wide intervals can extend beyond physical percentage bounds.')
p('Read transmissions, delay, reachability and RSU success together. Fewer frames alone are insufficient evidence of improvement if the protocol also fails to deliver. Connected-component reachability is an ideal graph opportunity and does not guarantee that the timed protocol over a lossy channel succeeds.')

h('17. Experiments, ablations and forged alerts')
p('Default requested populations are 50, 100, 200, 250, 300 and 500, with 30 run seeds per population: 180 mobility scenarios before per-scheme repetitions. Each scenario supplies identical positions, incident, RSUs, malicious set and trust values to every comparison. Channel random streams are reproducible but scheme-specific; they are not identical packet-by-packet contention draws.')
table(['Variant', 'What changes'], [('flooding / slotted1p / proposed', 'Headline comparison: same security, different forwarding rules.'), ('no_rad', 'Removes random assessment delay.'), ('no_slots', 'Removes distance-based slot ordering; priority tiers still contribute delay.'), ('no_coverage', 'Always forwards at timer expiry instead of checking uncovered neighbours.'), ('no_clustering', 'Treats all vehicles as noise; trusted vehicles become early bridges.'), ('no_cluster_forwarding', 'Keeps displayed roles but gives all vehicles the same forwarding tier.'), ('no_reputation', 'Removes trust gating for roles and endorsement acceptance.'), ('no_security', 'Removes proposal/voting/BLS and starts dissemination directly.'), ('ch_route', 'Restricts relay forwarding to a selected trusted-CH route and adds local quorum rounds.')])
p('Forgery tests compare full voting+BLS+reputation, voting+BLS without reputation, and no voting. Nearby malicious vehicles can sign the lie with their own keys, allowing the experiment to separate signer authenticity from event truth. Protection depends on trust separation and available colluders; the model does not demonstrate immunity to all malicious majorities or Sybil identities.')
p('The sensitivity command sweeps radio range across 150, 300 and 450 m, and malicious ratio across 0, 0.1, 0.2 and 0.3. EXP_ environment variables can override recognised configuration names; use scalar settings carefully because the generic type conversion is not a general parser for structured lists or dictionaries.')
p('Existing RESULTS.md identifies an older two-run-per-density result set. It must not be presented as the full current 30-seed experiment, and current code may have evolved since those files were generated. Re-run the intended configuration and inspect recorded experiment metadata before making final comparative claims.')

h('18. Web application and AI explanations')
p('The local web interface runs at http://127.0.0.1:8765 and reads saved experiment records. It lists simulation history, schemes and genuine/forged outcomes, with expandable recorded reports. It checks for new saved results about every five seconds. It is not a live radio-step telemetry feed.')
p('Users select a saved event and language, then request an explanation. Supported web languages include Tamil, Hindi, Telugu, Kannada, Malayalam, Bengali, Marathi, Gujarati, Punjabi and Odia. Terminal narration supports Tamil and English. Explanations are cached by evidence, scheme/event, language and model so matching requests can reuse text.')
p('narrator.py sends recorded evidence to a configured Groq model through an HTTPS API. The default model string in the code is qwen/qwen3.8-27b, overridable through GROQ_MODEL; actual service availability is not verified by this document. GROQ_API_KEY can be loaded server-side from the environment or local .env. Credentials are not included in the document.')
p('AI text explains evidence; it does not select clusters, produce metrics, verify signatures, prove an event or dispatch a response. The terminal provides explicit basic offline fallback for missing keys or failures. The web interface reports explanation errors rather than presenting untranslated English as a successful translation. Recorded values remain authoritative.')

h('19. Commands and practical usage')
table(['Command', 'Purpose'], [('python -m pip install -r requirements.txt', 'Install pinned project dependencies into the intended Python environment.'), ('python run.py', 'Open the default SUMO demonstration.'), ('python run.py demo', 'Run the GUI demonstration, normally for population 200, run 2.'), ('python run.py demo --scheme proposed ch_route', 'Choose the comparison sequence.'), ('python run.py single --all', 'Generate one scenario and print all scheme and forgery comparisons.'), ('python run.py single -d 100 -r 1 --debug', 'Choose population/run and include chronological detail.'), ('python run.py explain d200_r02 --scheme proposed', 'Explain saved forwarding decisions and generate supporting figures.'), ('python run.py experiment --densities 50 100 200 250 --runs 3', 'Small smoke experiment; insufficient alone for strong statistical conclusions.'), ('python run.py experiment', 'Run the configured complete multi-seed experiment.'), ('python run.py dashboard', 'Rebuild the static dashboard from existing outputs.'), ('python run.py sensitivity --density 200', 'Run configured range and malicious-population sweeps.'), ('python run.py web', 'Start the saved-results application on localhost.'), ('python run.py demo --narrate --language ta', 'Prepare Tamil narration before GUI playback.'), ('python run.py assistant d200_r02 --language en --offline', 'Read basic saved-run explanations without API calls.'), ('python run.py sync', 'Export local infrastructure data; upload only with configured endpoint.'), ('python run.py dispatch', 'Send queued events only to a configured HTTPS service endpoint.'), ('python -m pytest', 'Run validation tests; no full experiment is needed just to inspect this document.')])
p('For the local application, keep python run.py web in one terminal and run python run.py demo or single in another. Refresh after results have been saved. A graphical desktop is needed for SUMO-GUI; single/experiment/explain are alternatives for non-GUI use. Store API credentials outside source control and use test infrastructure endpoints for this simulation.')
p('Important outputs are raw/*.json.gz for detailed traces and timestamped history; processed/*.csv for per-run/aggregate metrics; summaries/ for narrative reports; figures/ for plots and screenshots; dashboard.html for static presentation; and infrastructure/ for SQLite and cloud export. Paths sit below results/urban by default.')

h('20. Recorded example, limitations and viva answers')
raw_path = ROOT / 'results/urban/raw/d200_r02.json.gz'
if raw_path.exists():
    with gzip.open(raw_path, 'rt', encoding='utf-8') as f: raw = json.load(f)
    sc = raw['scenario']
    p('The following values are read from the existing saved d200_r02 trace during document generation. They describe that saved run, not a fresh validation of current code or an average over many scenarios.')
    table(['Snapshot item', 'Recorded value'], [(k, sc.get(k, 'not recorded')) for k in ('run_id', 'seed', 'density', 'incident_time_s', 'source_id', 'incident_xy')])
    rows = []
    for name in ('flooding', 'slotted1p', 'proposed', 'ch_route'):
        if name in raw['protocols']:
            m = raw['protocols'][name]['metrics']
            rows.append((name, m.get('transmissions'), m.get('reachability'), m.get('rsu_received'), m.get('rsu_delay_ms')))
    table(['Scheme', 'All frames', 'Reachability fraction', 'RSU received', 'RSU delay ms'], rows)
p('The project assumes ideal sensor truth for honest voters, globally known trust, deterministic simulation identities, a fixed topology during dissemination and uniform radio range. It omits fading, building obstruction, realistic beacon interference, multiple concurrent incidents and actual backhaul latency. It therefore supports controlled algorithm comparison under stated assumptions, not a deployment safety guarantee.')
p('Future extensions include mobility during radio routing, channel fading and capture, noisy or correlated sensors, distributed reputation exchange, secure PKI with proof-of-possession enrolment, background traffic, congestion-aware routing, concurrent alerts and separately evaluated infrastructure-service reliability.')
table(['Viva question', 'Answer grounded in the implementation'], [('Why use clustering?', 'To give useful trusted representatives earlier relay opportunities and reduce redundant broadcasting.'), ('Why DBSCAN?', 'It detects dense spatial groups and sparse noise without fixing the cluster count.'), ('Why use both trust and signatures?', 'Trust estimates observed behaviour; signatures authenticate listed signers. Neither alone proves event truth.'), ('Why does the source need votes?', 'A valid source signature alone is insufficient corroboration of a claimed accident.'), ('Does BLS detect a fake crash?', 'No. It verifies signatures; honest observation and endorsement policy determine corroboration.'), ('Why use random delay?', 'To reduce simultaneous forwarding decisions among candidates with the same distance slot and tier.'), ('Why freeze vehicles?', 'The radio trace uses one incident snapshot and is slowed for presentation.'), ('Is the route shortest in time?', 'No. Dijkstra minimises summed link distance in metres among allowed nodes.'), ('Can an isolated DBSCAN vehicle help?', 'Yes. A trusted noise vehicle has early BRIDGE priority in proposed broadcasting.'), ('Does every cluster have exactly one head?', 'No. Greedy radio coverage can select several or find no eligible head.'), ('Does the project automatically call an ambulance?', 'No. Accepted CH-route records enter a local outbox unless a service endpoint is configured and dispatch is explicitly run.'), ('Is it machine learning?', 'DBSCAN is unsupervised clustering. There is no trained crash-classification or reinforcement-learning model in the decision pipeline.'), ('Is batch authentication faster pairing math?', 'No. It verifies per-event aggregates and deduplicates event IDs.'), ('What is the most defensible evaluation?', 'Repeatable multi-seed comparisons reporting delivery, delay, overhead and uncertainty together.')])

h('21. Configuration reference and source index')
p('The table below extracts top-level literal assignments from the current config.py. Runtime EXP_ overrides and the RSU junction-coordinate lookup can change effective settings. Comments that attribute parameter values to standards or papers require verification against original publications before thesis citation.')
tree = ast.parse((ROOT / 'config.py').read_text(encoding='utf-8'))
params = []
for node in tree.body:
    if not isinstance(node, ast.Assign): continue
    try: value = ast.literal_eval(node.value)
    except (ValueError, TypeError): continue
    for target in node.targets:
        if isinstance(target, ast.Name) and target.id.isupper():
            params.append((target.id, str(value)))
        elif isinstance(target, ast.Tuple) and isinstance(value, tuple):
            for target_item, part in zip(target.elts, value):
                if isinstance(target_item, ast.Name): params.append((target_item.id, str(part)))
table(['Configuration name', 'Literal default in source'], params)
p('Primary evidence for this explanation is the local source tree: PROJECT_GUIDE.md; config.py; scenario.py; clustering.py; corroboration.py; channel.py; protocols.py; infrastructure.py; rsu.py; experiment.py; metrics.py; demo.py; presentation.py; run.py; narrator.py; web_results.py; report.py; figures.py; network/build_urban.py; requirements.txt; and validation/*.py. Existing saved examples and screenshots are identified separately above.')
p('Algorithm names appearing in source comments include DBSCAN, Beta reputation, watchdog monitoring, greedy set cover, Dijkstra routing, slotted 1-persistence and BLS aggregate signatures. This document provides implementation explanations rather than a verified literature survey. Project novelty or paper-specific improvements need a separate comparison against the original research papers.')

# Apply explicit font attributes to every actual text run, including tables/header/footer.
all_paragraphs = list(doc.paragraphs) + list(sec.header.paragraphs) + list(sec.footer.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells: all_paragraphs.extend(cell.paragraphs)
for para in all_paragraphs:
    for run in para.runs:
        run.font.name = 'Gill Sans MT'
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'): fonts.set(qn('w:' + attr), 'Gill Sans MT')
doc.core_properties.title = 'Complete Project Explanation — Vehicular Emergency Alerts'
doc.core_properties.subject = 'Implementation, algorithms, formulas and SUMO simulation workflow'
doc.core_properties.author = 'Project Documentation'
doc.save(OUT)
print(OUT)
print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; images: {len(doc.inline_shapes)}')
