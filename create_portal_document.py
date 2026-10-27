from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / 'Project_Website_Documentation_Portal.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5)
styles['Normal'].font.color.rgb = RGBColor(45, 55, 68)
styles['Normal'].paragraph_format.space_after = Pt(4)
styles['Normal'].paragraph_format.line_spacing = 1.03
for nm, size in [('Title', 27), ('Heading 1', 17), ('Heading 2', 12)]:
    styles[nm].font.name = 'Aptos Display'; styles[nm].font.size = Pt(size)
    styles[nm].font.color.rgb = RGBColor(18, 72, 102)
    styles[nm].paragraph_format.space_before = Pt(7); styles[nm].paragraph_format.space_after = Pt(3)
header = sec.header.paragraphs[0]
header.text = 'PROJECT WEBSITE DOCUMENTATION  |  VEHICULAR EMERGENCY ALERTS'
header.style = styles['Caption']
footer = sec.footer.paragraphs[0]; footer.alignment = 2
footer.add_run('Project portal draft  •  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def title(t, sub=None):
    doc.add_heading(t, 0)
    if sub: doc.add_paragraph(sub)
def h(t): doc.add_heading(t, 2)
def p(t): doc.add_paragraph(t)
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def page(): doc.add_page_break()

# Page 1
title('Trust-Aware Emergency Alerts', 'Project Website / Documentation Portal • Five-page content draft')
p('A research prototype for carrying a verified accident alert from nearby vehicles to roadside infrastructure, including cases where the incident is outside direct roadside-unit radio range.')
h('1. Home Page')
p('The home page should introduce the project in one sentence, show the system flow, and provide clear links to the simulator, experiment dashboard, source repository, and report. A status panel can summarize the chosen scenario, vehicle count, and latest saved experiment. Treat results as research outputs, not live road-safety advice.')
p('Suggested hero: “Faster, more reliable vehicle-to-infrastructure emergency alerts through trusted relays and duplicate-broadcast suppression.” Primary actions: Explore the architecture • View results • Reproduce an experiment.')
h('2. Introduction and Background')
p('Connected vehicles can report hazards over vehicle-to-vehicle (V2V) links, while roadside units (RSUs) connect traffic to infrastructure. An incident beyond an RSU’s direct radio range needs intermediate vehicles to relay its message. Blind rebroadcasting can create redundant transmissions and channel contention; relying on one report can also expose the network to false alerts. This project studies a combined approach: organize nearby vehicles into clusters, prefer reputable relays, reduce redundant forwarding, and require signed peer endorsements before infrastructure accepts an alert.')
h('3. Problem Statement')
p('How can an accident alert reach an RSU through a mobile vehicular network with fewer unnecessary transmissions, while limiting acceptance of forged reports? The prototype evaluates delivery, delay, radio transmissions, congestion indicators, and forged-alert acceptance under controlled simulation assumptions. It does not detect real crashes or contact emergency services by itself.')

# Page 2
page(); h('4. Objectives and Research Gap')
bullets(['Reduce redundant alert broadcasts while retaining useful reachability.','Route alerts through reputable cluster representatives when direct RSU delivery is unavailable.','Require distinct signed endorsements and verify certificates at receiving vehicles and RSUs.','Compare forwarding strategies and component ablations under reproducible seeds.','Expose assumptions and limitations alongside performance results.'])
p('The research gap addressed by this prototype is the integration and evaluation of trust-aware cluster forwarding, broadcast suppression, and alert corroboration in one reproducible workflow. This is a project motivation, not a claim that the approach is novel or superior to all published systems. The source papers cited in implementation comments should be checked against original publications before publication.')
h('5. Proposed System and Architecture')
p('Pipeline: SUMO mobility → vehicle positions and neighbors → DBSCAN clusters → trust-gated representatives → incident and signed votes → radio forwarding → RSU certificate validation → local event store and dispatch outbox. A small local web application reads saved run data and presents metrics and explanations. Optional cloud export and emergency dispatch require separately configured HTTPS endpoints.')
p('Vehicles model V2V radio; two RSUs model infrastructure access. The RSU validates the certificate, registers accepted event IDs once, and stores accepted events and trust observations in SQLite. A durable outbox holds pending transfers. Credentials are kept in environment variables; the browser does not receive the AI narration API key.')
h('6. Methodology and Algorithms')
bullets(['DBSCAN groups vehicles using 40 m epsilon and a minimum of three vehicles. Representatives greedily cover cluster members; relative speed breaks coverage ties.','Beta reputation is estimated from simulated watchdog observations. A trust threshold of 0.7 gates prioritized relays. Silence alone does not imply misconduct.','The proposed broadcast strategy uses representative priority, distance slots, randomized delay, and neighbor-coverage suppression. Baselines include blind flooding and slotted 1-persistence.','The CH-route comparison uses Dijkstra over the source, reputable non-malicious heads, and available RSUs. Each forwarding head collects radio-delivered votes; three distinct endorsements including the head are required.','Receivers verify BLS signatures with py_ecc. Deterministic simulation identities are used; this is not a deployed vehicular public-key infrastructure.'])

# Page 3
page(); h('7. Implementation and Technologies')
p('The simulator and experiment pipeline are Python-based. SUMO supplies urban vehicle mobility; the project models radio exchange and forwarding in its Python simulation. SQLite stores accepted events and dispatch work. The included local results web application uses a lightweight Python server and HTML/JavaScript; it displays saved results rather than a live step-by-step feed. The code also supports optional multilingual AI explanations, with offline fallback.')
p('For a local run, install the dependencies in requirements.txt and SUMO, then follow PROJECT_GUIDE.md. Typical commands are python run.py demo, python run.py experiment, and python run.py web. The web interface is served locally at http://127.0.0.1:8765. Keep API credentials in environment variables and never commit them.')
h('8. Accident Simulation and Emergency Workflow')
bullets(['A controlled incident is injected at a selected vehicle, time, and location; it is a staged event, not a physics-based collision detector.','Nearby committee members respond to the incident’s simulated ground truth. Distinct signatures are aggregated into a certificate when quorum is reached.','The alert is disseminated using the selected forwarding scheme. Vehicles and RSUs reject invalid certificates and deduplicate event IDs.','When an RSU accepts the alert, it records the event and queues an outbox item. An operator can export JSON, sync to a configured ingestion endpoint, or dispatch to a configured test endpoint.','External endpoints must use HTTPS and honor the Idempotency-Key on retries. No real emergency-service endpoint is bundled or enabled by default.'])
h('9. Results and Performance Evaluation')
p('The checked-in results summary reports a quick comparison with two independent scenarios per density, not the full 30-seed configuration currently declared in config.py. In that snapshot, the proposed method used fewer frames than flooding at all shown densities; at 200 vehicles it reported 24.5 versus 181 mean frames, with similar vehicle reachability (100% for both in the sample). Across those samples, RSU delivery was 100% for all compared schemes. These small samples are preliminary and do not establish statistical superiority.')
p('Evaluate RSU success and delay separately from vehicle-wide reachability. Report frames on air, redundant alerts, reception losses, time to reach 90%, channel busy ratio, and suppression. Use identical seeds and snapshots across baselines, include confidence intervals and ablations, and clearly label quick runs as exploratory. Results can change as more runs are added.')

# Page 4
page(); h('10. Source Code and Reproducibility')
p('The repository contains the simulator modules, configuration, experiment runner, results viewer, validation checks, and saved urban results. config.py centralizes outcome-affecting parameters; experiment metadata is saved with generated results. The seed rule is BASE_SEED + 100 × density + run index. Preserve the config, seed, software versions, and raw output when reproducing a figure.')
bullets(['Quick dashboard: python run.py experiment --densities 50 100 200 250 --runs 3','Configured experiment: python run.py experiment','Refresh existing dashboard: python run.py dashboard','Local results portal: python run.py web','See PROJECT_GUIDE.md for full setup, options, and integration variables.'])
p('For a public portal, present source and data links, a commit identifier, environment/setup steps, exact commands, and a downloadable configuration snapshot. Do not expose .env or secrets. Add a clear version/date label to each result page.')
h('11. Research Papers and References')
p('Implementation comments identify the following sources or standards as parameter motivation. Verify bibliographic details and the original text before final academic submission; the repository notes that the source papers were not supplied for this documentation draft.')
bullets(['Wisitpongphan et al. (2007), broadcast storm mitigation in vehicular ad hoc networks; slotted 1-persistence baseline.','Marti et al. (2000), watchdog/pathrater mechanisms for detecting misbehaving nodes; motivation for watchdog-based observations.','Chen & Wu (2024), clustering parameter reference cited in config.py (full citation must be confirmed).','Trullols et al. (2010), RSU placement motivation cited in config.py (full citation must be confirmed).','ETSI EN 302 636-4-1 and ETSI TS 102 687; hop-limit and channel-busy-ratio conventions cited by the implementation.','SAE J2735; emergency-message payload scale cited as implementation context.'])
p('Add complete titles, venue/publisher, volume/pages, DOI or stable URL, and access date where applicable. Distinguish papers that inspired parameters from standards implemented or approximated by the simulator.')
h('12. Limitations and Future Work')
bullets(['The incident is staged; honest votes observe simulated ground truth and trust is shared.','Radio is a unit-disk approximation: decode, carrier-sense, and interference ranges are equal. SUMO is frozen during the incident dissemination snapshot.','Only two RSUs are modeled; real sensor errors, fading, backhaul latency, physical networking, cloud uptime, and operational dispatch are outside the radio model.','BLS identities are deterministic simulation identities, not a deployed PKI; batch authentication verifies each event individually.','Expand to larger independent samples, mobility during routing, realistic propagation, sensor uncertainty, privacy-preserving credentials, and hardware/network-in-the-loop evaluation.'])
p('Deployment direction: publish the static documentation portal and generated result assets on Vercel or another static host. Keep simulator execution, secrets, and optional API integrations on an appropriate backend; never imply the static portal is an operational emergency system.')

doc.save(OUT)
print(OUT)
